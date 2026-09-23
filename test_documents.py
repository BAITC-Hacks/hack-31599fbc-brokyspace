from __future__ import annotations

import io

import pandas as pd
import pytest
from docx import Document
from openpyxl import Workbook

from src.documents import CaseDocumentStore


def test_text_document_is_stored_deduplicated_and_searchable(tmp_path):
    store = CaseDocumentStore(tmp_path / "case_documents")
    known_gid = "100000003684369100"
    content = f"Служебная записка по клиенту {known_gid}. Контрагент Альфа.".encode("utf-8")

    record, created = store.add_document(
        "../проверка клиента.txt", content, known_gids={known_gid}
    )
    duplicate, created_again = store.add_document(
        "duplicate.txt", content, known_gids={known_gid}
    )

    assert created is True
    assert created_again is False
    assert duplicate.document_id == record.document_id
    assert record.filename == "проверка клиента.txt"
    assert record.detected_gids == [known_gid]
    assert len(store.list_documents()) == 1
    assert store.search("Альфа")[0]["filename"] == "проверка клиента.txt"
    assert store.search("проверка клиента")[0]["filename"] == "проверка клиента.txt"
    assert store.search(known_gid)[0]["detected_gids"] == [known_gid]


def test_unknown_gids_are_not_linked_to_graph(tmp_path):
    store = CaseDocumentStore(tmp_path)
    record, _ = store.add_document(
        "note.md",
        "GID 999999999999999999 and 100000003684369100".encode(),
        known_gids={"100000003684369100"},
    )
    assert record.detected_gids == ["100000003684369100"]


def test_json_original_is_not_overwritten_by_metadata(tmp_path):
    store = CaseDocumentStore(tmp_path)
    content = b'{"company": "Gamma", "risk": "review"}'
    record, _ = store.add_document("evidence.json", content)

    assert (tmp_path / f"{record.document_id}.json").read_bytes() == content
    assert (tmp_path / f"{record.document_id}.meta.json").is_file()
    assert store.search("Gamma")[0]["filename"] == "evidence.json"


def test_docx_text_and_tables_are_extracted(tmp_path):
    document = Document()
    document.add_paragraph("Проверка компании Бета")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "GID"
    table.cell(0, 1).text = "100000003684369100"
    stream = io.BytesIO()
    document.save(stream)

    store = CaseDocumentStore(tmp_path)
    record, created = store.add_document(
        "memo.docx", stream.getvalue(), known_gids={"100000003684369100"}
    )

    assert created is True
    assert record.text_chars > 20
    assert store.search("Бета")[0]["filename"] == "memo.docx"
    assert record.detected_gids == ["100000003684369100"]


def test_csv_xlsx_and_parquet_are_extracted(tmp_path):
    store = CaseDocumentStore(tmp_path)
    csv_record, _ = store.add_document(
        "переводы июля.csv",
        "gid,amount\n100000003684369100,25000".encode("utf-8"),
        known_gids={"100000003684369100"},
        mime_type="text/csv",
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Клиенты"
    sheet.append(["gid", "status"])
    sheet.append(["100000003684369100", "review"])
    xlsx_stream = io.BytesIO()
    workbook.save(xlsx_stream)
    xlsx_record, _ = store.add_document(
        "реестр клиентов.xlsx",
        xlsx_stream.getvalue(),
        known_gids={"100000003684369100"},
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    parquet_stream = io.BytesIO()
    pd.DataFrame({"gid": ["100000003684369100"], "amount": [50_000]}).to_parquet(
        parquet_stream, index=False
    )
    parquet_record, _ = store.add_document(
        "операции.parquet",
        parquet_stream.getvalue(),
        known_gids={"100000003684369100"},
        mime_type="application/octet-stream",
    )

    assert csv_record.detected_gids == ["100000003684369100"]
    assert xlsx_record.detected_gids == ["100000003684369100"]
    assert parquet_record.detected_gids == ["100000003684369100"]
    assert store.search("Клиенты")[0]["filename"] == "реестр клиентов.xlsx"


def test_mime_mismatch_and_corrupt_spreadsheet_are_rejected(tmp_path):
    store = CaseDocumentStore(tmp_path)
    with pytest.raises(ValueError, match="не соответствует расширению"):
        store.add_document("report.pdf", b"not-pdf", mime_type="text/plain")
    with pytest.raises(ValueError, match="Не удалось извлечь текст"):
        store.add_document(
            "broken.xlsx",
            b"not-an-excel-workbook",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


def test_document_can_be_opened_and_deleted_safely(tmp_path):
    store = CaseDocumentStore(tmp_path)
    record, _ = store.add_document("case note.txt", "Проверить клиента".encode("utf-8"))

    opened_record, original = store.get_original(record.document_id)
    assert opened_record.filename == "case note.txt"
    assert original.decode("utf-8") == "Проверить клиента"
    assert store.get_text(record.document_id) == "Проверить клиента"

    deleted = store.delete_document(record.document_id)
    assert deleted.document_id == record.document_id
    assert store.list_documents() == []
    with pytest.raises(ValueError, match="Некорректный идентификатор"):
        store.delete_document("../../outside")


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("malware.exe", b"not executable", "не поддерживается"),
        ("empty.txt", b"", "пуст"),
    ],
)
def test_invalid_documents_are_rejected(tmp_path, filename, content, message):
    with pytest.raises(ValueError, match=message):
        CaseDocumentStore(tmp_path).add_document(filename, content)
