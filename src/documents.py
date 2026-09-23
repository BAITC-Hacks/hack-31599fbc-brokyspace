from __future__ import annotations

import hashlib
import io
import json
import logging
import re
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd


ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".xlsx", ".xls", ".parquet"
}
MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
GID_PATTERN = re.compile(r"\b\d{12,20}\b")
DOCUMENT_ID_PATTERN = re.compile(r"^[0-9a-f]{16}$")
GENERIC_MIME_TYPES = {"", "application/octet-stream", "binary/octet-stream"}
MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip"},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "text/x-markdown"},
    ".csv": {"text/csv", "text/plain", "application/csv", "application/vnd.ms-excel"},
    ".json": {"application/json", "text/json", "text/plain"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/zip"},
    ".xls": {"application/vnd.ms-excel"},
    ".parquet": {"application/vnd.apache.parquet", "application/x-parquet"},
}
LOGGER = logging.getLogger(__name__)


@dataclass(frozen=True)
class DocumentRecord:
    document_id: str
    filename: str
    extension: str
    size_bytes: int
    sha256: str
    added_at: str
    text_chars: int
    detected_gids: list[str]
    mime_type: str = ""
    status: str = "Обработан"


class CaseDocumentStore:
    """Local-only store for user-supplied investigation documents."""

    def __init__(self, root: str | Path = "case_documents") -> None:
        self.root = Path(root)

    @staticmethod
    def _safe_filename(filename: str) -> str:
        name = Path(str(filename)).name.strip()
        name = re.sub(r"[^\w.()\- ]+", "_", name, flags=re.UNICODE).strip(" .")
        if not name:
            raise ValueError("У документа отсутствует допустимое имя")
        return name[:180]

    @staticmethod
    def _decode_text(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    @staticmethod
    def _validate_mime(extension: str, mime_type: str) -> str:
        normalized = str(mime_type or "").split(";", 1)[0].strip().lower()
        if normalized in GENERIC_MIME_TYPES:
            return normalized
        if normalized not in MIME_TYPES[extension]:
            raise ValueError(
                f"Тип файла {normalized!r} не соответствует расширению {extension}. "
                "Проверьте файл и повторите загрузку."
            )
        return normalized

    @classmethod
    def _extract_table(cls, extension: str, content: bytes) -> str:
        if extension == ".csv":
            decoded = cls._decode_text(content)
            frame = pd.read_csv(io.StringIO(decoded), sep=None, engine="python", nrows=10_000)
            if frame.empty and not len(frame.columns):
                raise ValueError("CSV не содержит таблицу")
            return frame.iloc[:, :100].to_csv(index=False)
        if extension in {".xlsx", ".xls"}:
            workbook = pd.ExcelFile(io.BytesIO(content))
            if not workbook.sheet_names:
                raise ValueError("В книге нет листов")
            blocks = []
            for sheet_name in workbook.sheet_names[:20]:
                frame = workbook.parse(sheet_name, nrows=10_000).iloc[:, :100]
                blocks.append(f"Лист: {sheet_name}\n{frame.to_csv(index=False)}")
            return "\n\n".join(blocks)
        if extension == ".parquet":
            frame = pd.read_parquet(io.BytesIO(content)).head(10_000).iloc[:, :100]
            if frame.empty and not len(frame.columns):
                raise ValueError("Parquet не содержит таблицу")
            return frame.to_csv(index=False)
        raise ValueError(f"Неподдерживаемый табличный формат: {extension}")

    @classmethod
    def _extract_text(cls, extension: str, content: bytes) -> str:
        if extension in {".txt", ".md"}:
            return cls._decode_text(content)
        if extension == ".json":
            payload = json.loads(cls._decode_text(content))
            return json.dumps(payload, ensure_ascii=False, indent=2)
        if extension in {".csv", ".xlsx", ".xls", ".parquet"}:
            return cls._extract_table(extension, content)
        if extension == ".docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(blocks)
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        raise ValueError(f"Неподдерживаемый формат: {extension}")

    def add_document(
        self,
        filename: str,
        content: bytes,
        *,
        known_gids: set[str] | None = None,
        mime_type: str = "",
    ) -> tuple[DocumentRecord, bool]:
        safe_name = self._safe_filename(filename)
        extension = Path(safe_name).suffix.lower()
        if extension not in ALLOWED_EXTENSIONS:
            raise ValueError(f"Формат {extension or 'без расширения'} не поддерживается")
        if not content:
            raise ValueError("Документ пуст")
        if len(content) > MAX_DOCUMENT_BYTES:
            raise ValueError("Размер документа превышает 15 MB")
        normalized_mime = self._validate_mime(extension, mime_type)

        digest = hashlib.sha256(content).hexdigest()
        document_id = digest[:16]
        metadata_path = self.root / f"{document_id}.meta.json"
        if metadata_path.exists():
            return DocumentRecord(**json.loads(metadata_path.read_text(encoding="utf-8"))), False

        try:
            text = self._extract_text(extension, content).strip()
        except Exception as exc:
            LOGGER.warning("Document extraction failed for %s: %s", safe_name, exc)
            raise ValueError(f"Не удалось извлечь текст из {safe_name}: {exc}") from exc
        if not text:
            raise ValueError("В документе не найден извлекаемый текст; сканированные PDF требуют OCR")

        mentioned = sorted(set(GID_PATTERN.findall(text)))
        if known_gids is not None:
            mentioned = [gid for gid in mentioned if gid in known_gids]
        record = DocumentRecord(
            document_id=document_id,
            filename=safe_name,
            extension=extension,
            size_bytes=len(content),
            sha256=digest,
            added_at=datetime.now(timezone.utc).isoformat(),
            text_chars=len(text),
            detected_gids=mentioned,
            mime_type=normalized_mime,
        )
        self.root.mkdir(parents=True, exist_ok=True)
        (self.root / f"{document_id}{extension}").write_bytes(content)
        (self.root / f"{document_id}.txt").write_text(text, encoding="utf-8")
        metadata_path.write_text(json.dumps(asdict(record), ensure_ascii=False, indent=2), encoding="utf-8")
        return record, True

    @staticmethod
    def _validate_document_id(document_id: str) -> str:
        normalized = str(document_id).lower()
        if not DOCUMENT_ID_PATTERN.fullmatch(normalized):
            raise ValueError("Некорректный идентификатор документа")
        return normalized

    def get_record(self, document_id: str) -> DocumentRecord:
        normalized = self._validate_document_id(document_id)
        metadata_path = self.root / f"{normalized}.meta.json"
        if not metadata_path.is_file():
            raise FileNotFoundError("Документ не найден")
        return DocumentRecord(**json.loads(metadata_path.read_text(encoding="utf-8")))

    def get_text(self, document_id: str) -> str:
        normalized = self._validate_document_id(document_id)
        text_path = self.root / f"{normalized}.txt"
        if not text_path.is_file():
            raise FileNotFoundError("Извлечённый текст документа не найден")
        return text_path.read_text(encoding="utf-8", errors="replace")

    def get_original(self, document_id: str) -> tuple[DocumentRecord, bytes]:
        record = self.get_record(document_id)
        original_path = self.root / f"{record.document_id}{record.extension}"
        if not original_path.is_file():
            raise FileNotFoundError("Оригинал документа не найден")
        return record, original_path.read_bytes()

    def delete_document(self, document_id: str) -> DocumentRecord:
        record = self.get_record(document_id)
        targets = (
            self.root / f"{record.document_id}{record.extension}",
            self.root / f"{record.document_id}.txt",
            self.root / f"{record.document_id}.meta.json",
        )
        for target in targets:
            if target.is_file():
                target.unlink()
        return record

    def list_documents(self) -> list[DocumentRecord]:
        if not self.root.exists():
            return []
        records = []
        for path in self.root.glob("*.meta.json"):
            try:
                records.append(DocumentRecord(**json.loads(path.read_text(encoding="utf-8"))))
            except (OSError, ValueError, TypeError):
                continue
        return sorted(records, key=lambda item: item.added_at, reverse=True)

    def search(self, query: str, limit: int = 10) -> list[dict[str, object]]:
        terms = [term.lower() for term in re.findall(r"[\w\-]{3,}", query, flags=re.UNICODE)]
        if not terms:
            return []
        results = []
        for record in self.list_documents():
            text_path = self.root / f"{record.document_id}.txt"
            if not text_path.exists():
                continue
            text = text_path.read_text(encoding="utf-8", errors="replace")
            lowered = text.lower()
            filename_lowered = record.filename.lower()
            positions = [lowered.find(term) for term in terms if lowered.find(term) >= 0]
            filename_matches = sum(term in filename_lowered for term in terms)
            if not positions and not filename_matches:
                continue
            position = min(positions) if positions else 0
            start = max(0, position - 180)
            end = min(len(text), position + 420)
            snippet = re.sub(r"\s+", " ", text[start:end]).strip()
            score = sum(lowered.count(term) for term in terms) + 2 * filename_matches
            results.append(
                {
                    "document_id": record.document_id,
                    "filename": record.filename,
                    "score": score,
                    "snippet": snippet,
                    "detected_gids": record.detected_gids,
                }
            )
        ordered = sorted(results, key=lambda item: (-int(item["score"]), str(item["filename"])))
        return ordered[: max(1, min(limit, 25))]
